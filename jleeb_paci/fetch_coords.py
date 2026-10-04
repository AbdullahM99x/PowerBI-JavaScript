"""Fill LAT/LONG in jleeb_paci.xlsx from a PACI ArcGIS REST layer.

Run this from a machine that can reach PACI (their servers refuse my sandbox).
Usage:  pip install openpyxl requests
        python fetch_coords.py "<layer URL>" [PACI_FIELD]

<layer URL> is the REST layer that holds PACI numbers as points/polygons, e.g.
  https://<paci-gis-host>/arcgis/rest/services/<folder>/<service>/MapServer/<layerId>
Find it in the browser dev-tools (Network tab) while searching a PACI number on
Kuwait Finder (kuwaitfinder.com / paci.gov.kw). PACI_FIELD defaults to PACI_NUMBER.
Polygons are reduced to their centroid. Output is WGS84 (outSR=4326).
NOTE: the endpoint/field names are not verified - adjust to what you see.
"""
import sys, requests
from openpyxl import load_workbook

url, field = sys.argv[1].rstrip("/"), (sys.argv[2] if len(sys.argv) > 2 else "PACI_NUMBER")
wb = load_workbook("jleeb_paci.xlsx"); ws = wb.active
s = requests.Session()
miss = []
for r in range(2, ws.max_row + 1):
    paci = str(ws.cell(r, 2).value)
    j = s.get(url + "/query", params={"where": f"{field}='{paci}'", "outFields": field,
              "returnGeometry": "true", "returnCentroid": "true", "outSR": 4326, "f": "json"}, timeout=30).json()
    f = (j.get("features") or [None])[0]
    if not f: miss.append(paci); continue
    g = f.get("centroid") or f["geometry"]
    if "rings" in g:  # fallback: mean of ring vertices
        pts = g["rings"][0]; g = {"x": sum(p[0] for p in pts) / len(pts), "y": sum(p[1] for p in pts) / len(pts)}
    ws.cell(r, 3).value, ws.cell(r, 4).value = round(g["y"], 7), round(g["x"], 7)
wb.save("jleeb_paci.xlsx"); print("not found:", miss)
