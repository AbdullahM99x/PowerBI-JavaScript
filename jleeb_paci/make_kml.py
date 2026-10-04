"""Build KML from jleeb_paci_plots.xlsx once LAT/LONG columns are filled. Rows without coordinates are skipped."""
import sys
from xml.sax.saxutils import escape
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "jleeb_paci_plots.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "jleeb_paci_plots.kml"
ws = load_workbook(src, read_only=True).active
pm, skipped = [], 0
for n, area, block, plot, paci, lat, lon, note in ws.iter_rows(min_row=2, values_only=True):
    if lat is None or lon is None:
        skipped += 1; continue
    pm.append(f"<Placemark><name>PACI {paci}</name><description>Plot {escape(str(plot))}, Block {block}, {area}"
              f" (#{n})</description><Point><coordinates>{float(lon)},{float(lat)},0</coordinates></Point></Placemark>")
with open(out, "w", encoding="utf-8") as f:
    f.write('<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
            '<name>Jleeb Al-Shuyoukh - demolition notice 2026/1727</name>' + "".join(pm) + "</Document></kml>")
print(f"{len(pm)} placemarks written, {skipped} rows skipped (no coordinates)")
