"""Build jleeb_paci.xlsx from data.txt (transcribed from Kuwait Alyoum issue 1811, pp. 40-42)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rows = [l.split() for l in open("data.txt", encoding="utf-8") if l.strip()]
dup = {p for p in {r[0] for r in rows} if sum(r[0] == p for r in rows) > 1}

wb = Workbook(); ws = wb.active; ws.title = "Jleeb PACI"
ws.append(["#", "الرقم الآلي (PACI)", "LAT", "LONG", "Area", "Block", "Plot", "Note"])
for i, (paci, plot) in enumerate(rows, 1):
    notes = []
    if len(paci) != 8: notes.append("PACI has %d digits as printed - verify" % len(paci))
    if paci in dup: notes.append("duplicate PACI in decree")
    ws.append([i, paci, None, None, "جليب الشيوخ", 1, plot, "; ".join(notes)])
    ws.cell(i + 1, 2).number_format = "@"
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center")
for col, w in zip("ABCDEFGH", [6, 20, 14, 14, 14, 8, 8, 45]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save("jleeb_paci.xlsx")
print(len(rows), "rows written")
