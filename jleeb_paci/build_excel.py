from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from data import ROWS, FLAGS

wb = Workbook(); ws = wb.active; ws.title = "Jleeb Demolition List"
ws.append(["#", "Area", "Block", "Plot (القسيمة)", "PACI (الرقم الآلي)", "LAT", "LONG", "Note"])
for i, (paci, plot) in enumerate(ROWS, 1):
    ws.append([i, "Jleeb Al-Shuyoukh", 1, plot, paci, None, None, FLAGS.get(i, "")])
    ws.cell(i + 1, 5).number_format = "@"
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center")
for col, w in zip("ABCDEFGH", (6, 20, 8, 14, 18, 12, 12, 60)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
wb.save("jleeb_paci_plots.xlsx")
